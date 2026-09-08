import io
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from sensefusion.common import Result
from sensefusion.diagnostics import augment, pareto_minima
from sensefusion.workbook import workbook_bytes


def test_group_and_range_diagnostics_keep_physical_units_and_counts():
    result = Result("Diagnostic fixture", settings={"application": "assayreport", "call_parameters": {}})
    result.tables["predictions"] = pd.DataFrame(
        {
            "sample_id": ["a1", "a2", "b1", "b2"],
            "group": ["a", "a", "b", "b"],
            "reference": [1.0, 3.0, 5.0, 7.0],
            "prediction": [2.0, 2.0, 7.0, 5.0],
        }
    )
    augment(result, pd.DataFrame())
    groups = result.tables["group_diagnostics"].set_index("group")
    assert groups.loc["a", "rmse"] == 1
    assert groups.loc["b", "rmse"] == 2
    assert groups.loc["a", "bias"] == groups.loc["b", "bias"] == 0
    assert result.tables["reference_range_diagnostics"].n.sum() == 4


def test_two_response_nondominance_retains_ties_and_is_scale_invariant():
    values = np.array([[0.0, 3.0], [1.0, 2.0], [2.0, 1.0], [3.0, 0.0], [2.0, 2.0], [1.0, 2.0]])
    expected = [True, True, True, True, False, True]
    assert pareto_minima(values).tolist() == expected
    assert pareto_minima(values * np.array([1e-12, 1e6])).tolist() == expected
    order = [4, 3, 5, 1, 0, 2]
    np.testing.assert_array_equal(pareto_minima(values[order]), np.array(expected)[order])


def test_paired_comparison_restricts_to_the_same_observations():
    result = Result("Diagnostic fixture", settings={"application": "sensefusion", "call_parameters": {}})
    result.tables["predictions"] = pd.DataFrame(
        {
            "sample_id": ["a", "b", "c", "a", "b"],
            "group": ["a", "b", "c", "a", "b"],
            "model": ["training_mean"] * 3 + ["single_A"] * 2,
            "reference": [1.0, 2.0, 3.0, 1.0, 2.0],
            "prediction": [0.0, 0.0, 0.0, 1.0, 3.0],
        }
    )
    augment(result, pd.DataFrame())
    comparison = result.tables["paired_baseline_comparison"].iloc[0]
    assert comparison.paired_rows == comparison.paired_groups == 2
    assert np.isclose(comparison.delta_group_rmse, np.sqrt(0.5) - np.sqrt(2.5))


def test_excel_companion_preserves_small_values_and_literal_text():
    result = Result("Workbook fixture", settings={"input_name": "=1+1"})
    result.tables["values"] = pd.DataFrame({"id": ["001", "=1+1"], "value": [1e-12, np.nan]})
    workbook = load_workbook(io.BytesIO(workbook_bytes(result)), data_only=False)
    assert workbook["Read me"]["B4"].data_type == "s"
    sheet = workbook["1 values"]
    assert sheet["A3"].data_type == "s"
    assert sheet["B2"].value == 1e-12
    assert sheet["B3"].value is None
    assert sheet["B2"].font.name == "Times New Roman"


def test_diagnostic_catalog_links_to_actual_exported_figures():
    from sensefusion.cli import example_config, example_context, example_path
    from sensefusion.common import read_table
    from sensefusion.core import analyse
    from sensefusion.reporting import files_for_result

    path = example_path()
    result = analyse(read_table(path), **(example_config(path.name) | example_context(path.name)))
    files = files_for_result(result)
    catalog = json.loads(files["figure_catalog.json"])
    assert len(catalog) >= 4
    for item in catalog:
        for extension in ["png", "pdf", "svg"]:
            assert len(files[item["stem"] + "." + extension]) > 1000
        assert item["stem"] + ".png" in files["extended_report.html"].decode("utf-8")
    assert "results.xlsx" in files
