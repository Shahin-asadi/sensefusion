"""Readable Excel companion to the complete CSV tables, with literal text cells."""

import io
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side


def workbook_bytes(result):
    """Export tables without formulas, private paths, or inferred precision."""
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Read me"
    cover.append(["Analysis", result.title])
    cover.append(["Evidence availability", result.status])
    cover.append(["Version", result.settings.get("software_version", "not recorded")])
    cover.append(["Input", result.settings.get("input_name", "in-memory table")])
    cover.append(
        [
            "Use",
            "Complete CSV tables and the run/configuration records remain the reproducibility source. No formulas or macros are used.",
        ]
    )
    cover.append(["Table", "Question / interpretation"])
    for item in result.settings.get("diagnostics", []):
        cover.append([item["table"], item["question"] + " " + item["interpretation"]])
    for index, (name, table) in enumerate(result.tables.items(), 1):
        sheet = workbook.create_sheet((str(index) + " " + name)[:31])
        sheet.append([str(c) for c in table.columns])
        for row in table.itertuples(index=False, name=None):
            converted = []
            for value in row:
                if hasattr(value, "item"):
                    value = value.item()
                if isinstance(value, float) and not math.isfinite(value):
                    value = None
                if isinstance(value, (list, tuple, dict)):
                    value = str(value)
                converted.append(value)
            sheet.append(converted)
        sheet.freeze_panes = "A2"
        if len(table.columns):
            sheet.auto_filter.ref = sheet.dimensions
    thin = Side(style="thin", color="000000")
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Times New Roman", size=10)
                cell.alignment = Alignment(vertical="top", horizontal="center", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = (
                        "0.000E+00" if 0 < abs(cell.value) < 0.0001 or abs(cell.value) >= 1e6 else "0.0000"
                    )
        for cell in sheet[1]:
            cell.font = Font(name="Times New Roman", size=12, bold=True)
            cell.border = Border(top=thin, bottom=thin)
        for cell in sheet[sheet.max_row]:
            cell.border = Border(bottom=thin)
        for column in sheet.columns:
            width = max((len(str(c.value)) for c in list(column)[:50] if c.value is not None), default=10)
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(14, width + 2))
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    cover.column_dimensions["B"].width = 90
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
