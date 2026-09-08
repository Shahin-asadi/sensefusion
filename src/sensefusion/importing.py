"""Rectangular, role-neutral input. Numeric conversion happens after column mapping."""

import csv
import io
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .common import InputError

DEFAULT_MAX_CELLS = 2_000_000
DEFAULT_MAX_BYTES = 50 * 1024 * 1024


def _bytes(source):
    if hasattr(source, "getvalue"):
        value = source.getvalue()
    elif hasattr(source, "read"):
        position = source.tell() if hasattr(source, "tell") else None
        value = source.read()
        if position is not None:
            source.seek(position)
    else:
        value = Path(source).read_bytes()
    return value.encode("utf-8") if isinstance(value, str) else value


def sheet_names(source):
    payload = _bytes(source)
    if len(payload) > DEFAULT_MAX_BYTES:
        raise InputError(
            "Workbook exceeds 50 MB. Choose a smaller table or use an explicit file-size limit in the Python importer."
        )
    try:
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        names = book.sheetnames
        book.close()
        return names
    except Exception as exc:
        raise InputError(
            "Workbook: could not open this XLSX. Save an unencrypted .xlsx copy and select its data sheet."
        ) from exc


def _empty(value):
    return (
        value is None
        or (isinstance(value, str) and not value.strip())
        or (isinstance(value, float) and np.isnan(value))
    )


def inspect_source(
    source,
    filename=None,
    *,
    sheet=None,
    header_row=1,
    separator="auto",
    encoding="utf-8-sig",
    max_cells=DEFAULT_MAX_CELLS,
    max_bytes=DEFAULT_MAX_BYTES,
):
    if type(header_row) is not int or header_row < 1:
        raise InputError("Header row: enter a positive row number, starting at 1.")
    if type(max_cells) is not int or max_cells < 1 or type(max_bytes) is not int or max_bytes < 1:
        raise InputError("Import limits: cell and byte budgets must be positive integers.")
    payload = _bytes(source)
    if len(payload) > max_bytes:
        raise InputError(
            f"File has {len(payload):,} bytes, above the {max_bytes:,}-byte import limit. Select a smaller table or deliberately increase the import limit."
        )
    name = str(filename or getattr(source, "name", source)).lower()
    notices, selected_sheet = [], None
    if name.endswith(".xlsx"):
        try:
            cached = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            formulas = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
            choices = cached.sheetnames
            if sheet is not None and sheet not in choices:
                raise InputError(f"Worksheet {sheet!r} does not exist. Select one of: {', '.join(choices)}.")
            if sheet is None:
                selected_sheet = next((s for s in choices if s.strip().lower() == "data"), None)
                if selected_sheet is None:
                    selected_sheet = next(
                        (s for s in choices if cached[s].max_column >= 2 and cached[s].max_row > header_row), choices[0]
                    )
                notices.append(f"Worksheet selected: {selected_sheet}. Confirm this is the intended data sheet.")
            else:
                selected_sheet = sheet
            ws, fs = cached[selected_sheet], formulas[selected_sheet]
            if ws.max_row * ws.max_column > max_cells:
                raise InputError(
                    f"Worksheet {selected_sheet!r} declares {ws.max_row * ws.max_column:,} cells, above the {max_cells:,}-cell limit. Clear distant formatting or choose a smaller rectangular table."
                )
            rows, uncached = [], []
            for row_index, (values, expressions) in enumerate(zip(ws.iter_rows(), fs.iter_rows()), 1):
                rows.append([c.value for c in values])
                if row_index > header_row:
                    for value, expression in zip(values, expressions):
                        if expression.data_type == "f" and value.value is None:
                            uncached.append(expression.coordinate)
            cached.close()
            formulas.close()
            if uncached:
                raise InputError(
                    f"Worksheet {selected_sheet!r}: formulas at {', '.join(uncached[:3])} have no saved numeric result. Recalculate and save the workbook in Excel or LibreOffice, then upload again."
                )
        except InputError:
            raise
        except Exception as exc:
            raise InputError(
                "Workbook could not be read. Choose a data worksheet in an unencrypted XLSX and check its header row."
            ) from exc
        used_separator = None
    else:
        try:
            text = payload.decode(encoding)
        except (UnicodeDecodeError, LookupError) as exc:
            raise InputError(
                "Text encoding: select the file encoding in Import options, or save as UTF-8 CSV."
            ) from exc
        if separator == "auto":
            try:
                used_separator = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
            except csv.Error:
                used_separator = ","
        elif separator in [",", ";", "\t"]:
            used_separator = separator
        else:
            raise InputError("Separator: select comma, semicolon or tab.")
        try:
            rows = []
            for index, row in enumerate(csv.reader(io.StringIO(text), delimiter=used_separator, strict=True), 1):
                rows.append(row)
                if len(row) * index > max_cells:
                    raise InputError(
                        f"Text table exceeds the {max_cells:,}-cell limit. Choose a smaller table or increase the import limit explicitly."
                    )
        except csv.Error as exc:
            raise InputError(
                "Delimited text: a quoted cell or row is malformed. Check the separator and quote marks near the header."
            ) from exc
    if header_row > len(rows):
        raise InputError("Header row is beyond the available table. Select the row containing column names.")
    rows = rows[header_row - 1 :]
    removed_rows = 0
    while rows and all(_empty(v) for v in rows[-1]):
        rows.pop()
        removed_rows += 1
    if not rows:
        raise InputError("The selected sheet/table is empty. Select the data worksheet or a different header row.")
    width = max(map(len, rows))
    rows = [row + [None] * (width - len(row)) for row in rows]
    removed_columns = 0
    while width and all(_empty(row[width - 1]) for row in rows):
        width -= 1
        removed_columns += 1
    rows = [r[:width] for r in rows]
    if removed_rows or removed_columns:
        notices.append(
            f"Ignored {removed_rows} wholly empty trailing rows and {removed_columns} wholly empty trailing columns."
        )
    if not width:
        raise InputError("No populated columns remain. Select the correct data header.")
    if any(len(r) > len(rows[0]) for r in rows):
        raise InputError("Rows have more fields than the header. Check the separator or select the correct header row.")
    return rows, {
        "worksheet": selected_sheet,
        "header_row": header_row,
        "separator": used_separator,
        "encoding": encoding,
        "notices": notices,
    }


def read_table(
    source,
    filename=None,
    *,
    sheet=None,
    header_row=1,
    separator="auto",
    decimal=".",
    encoding="utf-8-sig",
    missing_tokens=(),
    header_overrides=None,
    max_cells=DEFAULT_MAX_CELLS,
    max_bytes=DEFAULT_MAX_BYTES,
    allow_empty=False,
):
    if decimal not in [".", ","]:
        raise InputError("Decimal mark: select a period or comma.")
    if not isinstance(missing_tokens, (list, tuple)) or any(not isinstance(t, str) for t in missing_tokens):
        raise InputError("Missing tokens: provide a list of text markers used in numeric columns.")
    rows, info = inspect_source(
        source,
        filename,
        sheet=sheet,
        header_row=header_row,
        separator=separator,
        encoding=encoding,
        max_cells=max_cells,
        max_bytes=max_bytes,
    )
    headers = [str(v).strip() if not _empty(v) else "" for v in rows[0]]
    original = headers.copy()
    if header_overrides is not None:
        if not isinstance(header_overrides, (list, tuple)) or len(header_overrides) != len(headers):
            raise InputError("Column names: provide exactly one name for each preview column.")
        headers = [str(v).strip() for v in header_overrides]
    blank = [str(i + 1) for i, h in enumerate(headers) if not h]
    if blank:
        raise InputError(
            f"Header: populated unnamed column(s) {', '.join(blank[:3])}. Set their names in Rename columns, or select the correct header row."
        )
    duplicates = pd.Index(headers)[pd.Index(headers).duplicated()].tolist()
    if duplicates:
        raise InputError(
            f"Header: duplicate column(s) {duplicates[:3]}. Column names must be unique; edit them in Rename columns."
        )
    if len(rows) < 2 and not allow_empty:
        raise InputError("The table has headers but no observations. Add your measurements below the header.")
    frame = pd.DataFrame(rows[1:], columns=headers, dtype=object)
    frame = frame.map(lambda v: np.nan if _empty(v) else v)
    info.update(
        decimal=decimal,
        missing_tokens=list(missing_tokens),
        column_mapping=[
            {"position": i + 1, "source_header": original[i], "analysis_header": name} for i, name in enumerate(headers)
        ],
        max_cells=max_cells,
        max_bytes=max_bytes,
    )
    frame.attrs["import_info"] = info
    return frame


def numeric_columns(frame, columns):
    missing = sorted(set(columns) - set(frame.columns))
    if missing:
        raise InputError("Missing columns: " + ", ".join(missing) + ". Select existing numeric columns.")
    if len(columns) != len(set(columns)):
        raise InputError("Numeric column selections must be unique; remove duplicate columns.")
    info = frame.attrs.get("import_info", {})
    tokens = set(info.get("missing_tokens", []))
    decimal = info.get("decimal", ".")
    raw = frame[columns].copy()

    def normalize(value):
        if _empty(value):
            return np.nan
        if isinstance(value, str):
            value = value.strip()
            if value in tokens:
                return np.nan
            if decimal == ",":
                value = value.replace(",", ".")
        return value

    normalized = raw.map(normalize)
    converted = normalized.apply(pd.to_numeric, errors="coerce")
    invalid = normalized.notna() & converted.isna()
    infinite = pd.DataFrame(np.isinf(converted.to_numpy(dtype=float)), index=frame.index, columns=columns)
    if invalid.any().any() or infinite.any().any():
        problems = []
        for row, col in np.argwhere((invalid | infinite).to_numpy())[:3]:
            problems.append(f"row {row + info.get('header_row', 1) + 1}, {columns[col]!r}: {raw.iloc[row, col]!r}")
        raise InputError(
            "Numeric values: "
            + "; ".join(problems)
            + ". Correct these cells, select the decimal mark, or declare the missing token in Import options. Infinite values are not valid measurements."
        )
    return converted.astype(float)
