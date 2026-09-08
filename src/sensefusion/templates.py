"""Small artificial input examples. Never use these rows as experimental evidence."""

import io
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side


def schema():
    slug = json.loads(Path(__file__).with_name("project.json").read_text())["slug"]
    if slug == "sensefusion":
        data = {
            "Reference": [2, 4, 6, 8, 10, 12, 14, 16],
            "Sensor A": [1.2, 2.1, 2.8, 4.3, 5.1, 5.7, 7.2, 7.9],
            "Sensor B": [3.1, 4.8, 7.2, 8.9, 11.3, 12.8, 15.0, 17.2],
        }
        config = {
            "target": "Reference",
            "id_col": None,
            "group_col": None,
            "block_columns": {"Sensor A": ["Sensor A"], "Sensor B": ["Sensor B"]},
        }
        roles = [
            ("Reference response", "Reference"),
            ("Measurement block A", "Sensor A"),
            ("Measurement block B", "Sensor B"),
        ]
        row_meaning = (
            "One artificial observation. Add a physical specimen column if your real file contains repeated readings."
        )
    elif slug == "calibshift":
        data = {"Specimen": [], "Condition": [], "Reference": [], "Signal": []}
        for i in range(8):
            for condition in ["Source", "Changed"]:
                data["Specimen"].append(f"{i + 1:03d}")
                data["Condition"].append(condition)
                data["Reference"].append(float((i + 1) * 5))
                data["Signal"].append((i + 1) * 4 + (0.3 if i % 2 else -0.2) + (2 if condition == "Changed" else 0))
        config = {
            "target": "Reference",
            "id_col": None,
            "group_col": "Specimen",
            "domain_col": "Condition",
            "source": "Source",
            "destinations": ["Changed"],
            "feature_columns": ["Signal"],
            "budgets": [0, 1, 2, 4],
        }
        roles = [
            ("Physical specimen / formulation", "Specimen"),
            ("Acquisition condition", "Condition"),
            ("Reference response", "Reference"),
            ("Measurement", "Signal"),
        ]
        row_meaning = (
            "One specimen measured in one condition. Specimen matches the same unchanged formulation across conditions."
        )
    elif slug == "matcheddoe":
        data = {"Time (min)": list(range(9)), "Response": [2.1, 3.8, 6.2, 7.9, 10.3, 11.7, 14.1, 15.9, 18.2]}
        config = {
            "factors": ["Time (min)"],
            "responses": ["Response"],
            "id_col": None,
            "model": "linear",
            "independent_runs": False,
        }
        roles = [("Numeric factor", "Time (min)"), ("Response", "Response")]
        row_meaning = (
            "One artificial run. Independence must be established from the real experimental design before inference."
        )
    else:
        data = {"Reference": [0, 2, 4, 6, 8, 10, 12, 14], "Comparison": [0.1, 2.2, 3.8, 6.1, 8.3, 9.8, 12.1, 14.2]}
        config = {
            "target": "Reference",
            "prediction_column": "Comparison",
            "id_col": None,
            "group_col": None,
            "feature_columns": [],
            "comparison_kind": "measured",
            "agreement_unit": "individual",
            "independent_pairs": False,
        }
        roles = [("Reference value", "Reference"), ("Measured or predicted comparison", "Comparison")]
        row_meaning = "One paired comparison. Group repeated readings by physical specimen when needed; confirm independence separately."
    return {
        "slug": slug,
        "data": data,
        "config": config,
        "roles": roles,
        "row_meaning": row_meaning,
        "artificial": True,
    }


def starter_frame(blank=False):
    frame = pd.DataFrame(schema()["data"])
    return frame.iloc[:0].copy() if blank else frame


def starter_bytes(format="csv", blank=False):
    spec = schema()
    frame = starter_frame(blank)
    if format == "csv":
        return frame.to_csv(index=False, lineterminator="\n").encode("utf-8-sig")
    if format != "xlsx":
        raise ValueError("Choose csv or xlsx.")
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append(list(row))
    rules = Side(style="thin", color="000000")
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row == 1:
                cell.font = Font(name="Times New Roman", size=12, bold=True)
                cell.border = Border(top=rules, bottom=rules)
            elif cell.row == sheet.max_row:
                cell.border = Border(bottom=rules)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = max(17, min(38, len(str(column[0].value)) + 5))
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    instructions = book.create_sheet("Instructions")
    notes = [
        "ARTIFICIAL SOFTWARE EXAMPLE — replace every example row before analysing real measurements.",
        spec["row_meaning"],
        "This row count is illustrative, not a minimum sample size or a recommended experimental design.",
        "Data starts in row 1 of the Data sheet. Add/remove rows and measurement columns freely.",
        "Templates are optional. Your own rectangular CSV/XLSX is supported through column mapping.",
        "Keep reference and comparison units compatible. Blank cells are missing; zero remains a value.",
        "No macros, formulas, protected cells, merged Data cells or fixed entry area are used.",
    ]
    for note in notes:
        instructions.append([note])
    instructions.append(["Column-role example"])
    for role, column in spec["roles"]:
        instructions.append([role, column])
    for row in instructions:
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10, bold=cell.row in [1, 8])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
        instructions.row_dimensions[row[0].row].height = 42 if row[0].row <= 7 else 26
    instructions.column_dimensions["A"].width = 90
    instructions.column_dimensions["B"].width = 30
    instructions.sheet_view.showGridLines = False
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
